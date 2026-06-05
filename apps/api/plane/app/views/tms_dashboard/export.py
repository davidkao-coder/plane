# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.
#
# TMS report export to Excel – Phase 2 / B4.

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from plane.app.permissions import ROLE, allow_permission
from plane.db.models import Issue, Project
from plane.utils.tms_health import is_issue_overdue

from .. import BaseAPIView


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


class TMSExportEndpoint(BaseAPIView):
    """GET /workspaces/<slug>/projects/<pid>/export-report/
    Streams an .xlsx workbook: one summary sheet + one full work-item sheet
    (with stage / feature / hours / overdue flag).
    """

    @allow_permission([ROLE.ADMIN, ROLE.MEMBER, ROLE.GUEST])
    def get(self, request, slug, project_id):
        try:
            project = Project.objects.get(pk=project_id, workspace__slug=slug)
        except Project.DoesNotExist:
            return Response_404()

        today = timezone.now().date()

        issues = (
            Issue.issue_objects.filter(project_id=project_id, workspace__slug=slug)
            .select_related(
                "state",
                "stage",
                "process_step",
                "feature",
                "feature__requirement",
                "feature__requirement__module",
            )
            .order_by("stage__sort_order", "process_step__sort_order", "sequence_id")
        )

        wb = Workbook()

        # ── Summary sheet ────────────────────────────────────────────────
        ws1 = wb.active
        ws1.title = "摘要"
        total = issues.count()
        completed = sum(1 for i in issues if i.state and i.state.group in ("completed", "cancelled"))
        overdue = sum(
            1 for i in issues if is_issue_overdue(i.target_date, i.state.group if i.state else None, today)
        )
        rows = [
            ("專案", project.name),
            ("代號", project.identifier),
            ("客戶", project.client_name or "—"),
            ("合約編號", project.contract_no or "—"),
            ("產出日期", today.isoformat()),
            ("工作項目總數", total),
            ("已完成", completed),
            ("逾期", overdue),
            ("完成率", f"{round(completed / total * 100, 1) if total else 0}%"),
        ]
        ws1["A1"] = "項目"
        ws1["B1"] = "內容"
        _style_header(ws1, 2)
        for idx, (k, v) in enumerate(rows, start=2):
            ws1.cell(row=idx, column=1, value=k).font = Font(bold=True)
            ws1.cell(row=idx, column=2, value=v)
        ws1.column_dimensions["A"].width = 18
        ws1.column_dimensions["B"].width = 40

        # ── Work items sheet ─────────────────────────────────────────────
        ws2 = wb.create_sheet("工作項目")
        headers = [
            "編號", "分類", "需求", "功能", "階段", "工序", "名稱",
            "狀態", "預估工時", "實際工時", "開始日", "截止日", "逾期",
        ]
        ws2.append(headers)
        _style_header(ws2, len(headers))
        for i in issues:
            feat = i.feature
            req = feat.requirement if feat else None
            mod = req.module if req else None
            ws2.append([
                i.sequence_id,
                mod.name if mod else "",
                req.requirement_id if req else "",
                feat.feature_id if feat else "",
                i.stage.name if i.stage else "",
                i.process_step.name if i.process_step else "",
                i.name,
                i.state.name if i.state else "",
                float(i.estimate_hours) if i.estimate_hours is not None else None,
                float(i.actual_hours) if i.actual_hours is not None else None,
                i.start_date.isoformat() if i.start_date else "",
                i.target_date.isoformat() if i.target_date else "",
                "是" if is_issue_overdue(i.target_date, i.state.group if i.state else None, today) else "",
            ])
        widths = [8, 14, 10, 10, 12, 12, 40, 12, 10, 10, 12, 12, 6]
        for col, w in enumerate(widths, start=1):
            ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = w
        ws2.freeze_panes = "A2"

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"{project.identifier}_report_{today.isoformat()}.xlsx"
        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp


def Response_404():
    from rest_framework.response import Response
    from rest_framework import status

    return Response({"error": "Project not found"}, status=status.HTTP_404_NOT_FOUND)
